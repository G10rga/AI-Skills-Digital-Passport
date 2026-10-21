"""Admin Excel upload: parse BTU workbook layout, ingest, sync, issue."""
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)

HEADERS = ["Student ID", "First Name", "Last Name", "BTU Email (Test)", "Program",
           "Education Level", "Passport ID (Existing)", "First Issued (Existing)",
           "Student Status", "Record ID", "Skill ID", "Skill Name", "Skill Category",
           "Badge/Stamp Code", "Skill Description", "Date Earned",
           "Passport Sync Status", "Synced Passport Version", "Synced At (UTC)",
           "Sync Error / Notes"]


@pytest.fixture(autouse=True)
def fresh():
    assert client.post("/admin/reset").status_code == 200
    yield


def workbook(rows, header_row=4):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=header_row, column=1, value=HEADERS[0])
    for j, h in enumerate(HEADERS, start=1):
        ws.cell(row=header_row, column=j, value=h)
    for i, row in enumerate(rows, start=header_row + 1):
        for j, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(row=i, column=j, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def new_row(record_id="REC-9200", skill_id="SKL-099", sync_status=None):
    return ["BTU-T-0001", "Giorgi", "Kapanadze", "btu-t-0001@example.org",
            "Business Administration (Management)", "Bachelor's", "BTU-AI-000125",
            "2026-09-15", "Active", record_id, skill_id, "Excel Imported Skill",
            "AI Skill", None, "Imported from an uploaded workbook.", "2026-10-20",
            sync_status, None, None, None]


def upload(data, filename="import.xlsx", **form):
    return client.post("/admin/records/upload", data=form,
                       files={"file": (filename, data,
                                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def test_upload_btu_layout_ingests_and_syncs_new_version():
    before = client.get("/passports/BTU-AI-000125").json()
    r = upload(workbook([new_row()]))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["rows"] == 1 and body["imported"] == 1 and body["skipped"] == 0
    assert body["sync"]["synced"] >= 1
    after = client.get("/passports/BTU-AI-000125").json()
    assert len(after["versions"]) == len(before["versions"]) + 1
    skill = next(s for s in after["versions"][-1]["skills"] if s["recordId"] == "REC-9200")
    assert skill["skillName"] == "Excel Imported Skill" and skill["badgeCode"] == "BDG-AI-001"


def test_upload_existing_record_updates_without_duplicating():
    assert upload(workbook([new_row()])).status_code == 200
    versions = len(client.get("/passports/BTU-AI-000125").json()["versions"])
    r = upload(workbook([new_row()]))
    assert r.json()["imported"] == 0 and r.json()["updated"] == 1
    assert len(client.get("/passports/BTU-AI-000125").json()["versions"]) == versions


def test_upload_without_sync_only_ingests():
    r = client.post("/admin/records/upload", data={"sync": "false"},
                    files={"file": ("i.xlsx", workbook([new_row("REC-9201", "SKL-098")]),
                                    "application/vnd.ms-excel")})
    assert r.status_code == 200 and r.json()["sync"] is None
    rec = next(x for x in client.get("/records").json() if x["recordId"] == "REC-9201")
    assert rec["syncStatus"] is None


def auth(email):
    code = client.post("/auth/request-code", json={"email": email}).json()["devCode"]
    tok = client.post("/auth/verify-code", json={"email": email, "code": code}).json()["token"]
    return {"Authorization": f"Bearer {tok}"}


def test_upload_new_student_without_passport_id_gets_passport_automatically():
    # the reported case: brand-new student, Passport ID + First Issued blank,
    # row already marked SYNCED in the uploaded file
    row = ["BTU-T-06701", "Lasha", "Giorgadze", "lasha.giorgadze.2@btu.edu.ge",
           "Computer science and artificial inteligence", "Bachelor's", None, None,
           "Active", "REC-06701", "SKL-001", "AI Literacy", "AI Skill",
           '=IF(M2="AI Skill","BDG-AI-001","BDG-AI-003")',
           "Fundamentals of AI capabilities, limitations, and responsible use.",
           "2026-09-18", "SYNCED", "V1", "2026-10-17 10:00 UTC", None]
    r = upload(workbook([row]))
    assert r.status_code == 200, r.text
    assert r.json()["newStudents"] == 1
    summaries = {p["studentId"]: p for p in client.get("/passports").json()}
    assert "BTU-T-06701" in summaries and summaries["BTU-T-06701"]["latestVersion"] == 1
    # the student can immediately claim it — no staff step needed
    h = auth("lasha.giorgadze.2@btu.edu.ge")
    req = client.post("/passport-requests", headers=h, json={})
    assert req.status_code == 201 and req.json()["status"] == "READY"
    mine = client.get("/me/passport", headers=h).json()
    assert mine["studentId"] == "BTU-T-06701" and len(mine["versions"]) == 1
    assert mine["versions"][0]["skills"][0]["skillName"] == "AI Literacy"


def test_passport_request_heals_record_imported_without_passport_id():
    # data that reached the DB before auto-assignment existed (blank ID)
    from app.main import store
    store.upsert_record({"studentId": "BTU-T-7777", "firstName": "Ana", "lastName": "K",
                         "email": "ana.k@example.org", "passportId": "", "firstIssued": "",
                         "recordId": "REC-7777", "skillId": "SKL-001", "skillName": "AI Literacy",
                         "skillCategory": "AI Skill", "badgeCode": "BDG-AI-001",
                         "skillDescription": "Fundamentals", "dateEarned": "2026-10-01",
                         "syncStatus": "SYNCED"})
    h = auth("ana.k@example.org")
    req = client.post("/passport-requests", headers=h, json={})
    assert req.status_code == 201, req.text
    assert req.json()["status"] == "READY" and req.json()["passportId"].startswith("BTU-AI-")
    assert client.get("/me/passport", headers=h).json()["versions"][0]["skills"]


def test_upload_rejects_non_xlsx_and_garbage():
    assert upload(b"a,b,c", filename="data.csv").status_code == 400
    assert upload(b"not a workbook").status_code == 400


def test_upload_camelcase_header_row_1():
    wb = openpyxl.Workbook()
    ws = wb.active
    keys = ["studentId", "firstName", "lastName", "email", "program", "educationLevel",
            "passportId", "firstIssued", "studentStatus", "recordId", "skillId",
            "skillName", "skillCategory", "badgeCode", "skillDescription", "dateEarned"]
    for j, k in enumerate(keys, start=1):
        ws.cell(row=1, column=j, value=k)
    for j, v in enumerate(new_row("REC-9202", "SKL-097")[:16], start=1):
        ws.cell(row=2, column=j, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    r = upload(buf.getvalue())
    assert r.status_code == 200, r.text
    assert r.json()["imported"] == 1
