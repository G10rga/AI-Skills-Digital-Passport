"""Sheet sources. Google Sheets is the system of record for input columns
A-P; columns Q-T are system-managed and written back after a sync.

All sources return rows as SkillRecord-shaped dicts (camelCase). Row
tracking (`_sheetRow`, 1-based sheet row) lets write-back hit the exact
Q-T cells of the row that produced a record.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Optional, Protocol

from .store import load_seed

# Sheet columns A-T, in order (see the hackathon test workbook / records.json).
COLUMNS = ["studentId", "firstName", "lastName", "email", "program", "educationLevel",
           "passportId", "firstIssued", "studentStatus", "recordId", "skillId", "skillName",
           "skillCategory", "badgeCode", "skillDescription", "dateEarned",
           "syncStatus", "syncedVersion", "syncedAt", "syncNotes"]


class SheetSource(Protocol):
    def fetch(self) -> list[dict]: ...
    def write_back(self, updates: list[dict]) -> None: ...


NULLABLE = ("syncStatus", "syncedVersion", "syncedAt", "syncNotes")


def _clean(rows: list[dict]) -> list[dict]:
    for r in rows:
        for k in COLUMNS:
            v = r.get(k)
            v = v.strip() if isinstance(v, str) else ("" if v is None else v)
            r[k] = (v or None) if k in NULLABLE else v
    return [r for r in rows if r.get("recordId")]


class LocalJsonSheet:
    """Seed JSON (frontend/src/data/records.json) — the offline demo sheet."""

    def __init__(self, path: Path, writeback: bool = False):
        self.path, self.writeback_enabled = path, writeback

    def fetch(self) -> list[dict]:
        return _clean(load_seed(self.path))

    def write_back(self, updates: list[dict]) -> None:
        if not self.writeback_enabled:
            return
        rows = {r["recordId"]: r for r in load_seed(self.path)}
        for u in updates:
            if u["recordId"] in rows:
                rows[u["recordId"]].update({k: u[k] for k in
                    ("syncStatus", "syncedVersion", "syncedAt", "syncNotes")})
        self.path.write_text(json.dumps(list(rows.values()), indent=1, ensure_ascii=False),
                             encoding="utf-8")


class XlsxSheet:
    """The hackathon test workbook (.xlsx). Header is on row 4, data from row 5."""

    def __init__(self, path: Path, writeback: bool = False):
        self.path, self.writeback_enabled = path, writeback

    def _book(self):
        try:
            import openpyxl  # optional: only for xlsx deployments
        except ImportError as e:
            raise RuntimeError("SHEET_SOURCE=xlsx requires `pip install openpyxl`") from e
        return openpyxl.load_workbook(self.path)

    def fetch(self) -> list[dict]:
        ws = self._book().active
        rows = []
        for i, values in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
            if not values or not any(values):
                continue
            r = {k: ("" if v is None else str(v)) for k, v in zip(COLUMNS, values)}
            r["_sheetRow"] = i
            rows.append(r)
        return _clean(rows)

    def write_back(self, updates: list[dict]) -> None:
        if not self.writeback_enabled:
            return
        wb = self._book()
        ws = wb.active
        by_id = {u["recordId"]: u for u in updates}
        for i, values in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
            u = by_id.get(str(values[9]) if values and values[9] else "")
            if u:
                for col, key in ((17, "syncStatus"), (18, "syncedVersion"),
                                 (19, "syncedAt"), (20, "syncNotes")):
                    ws.cell(row=i, column=col, value=u[key])
        wb.save(self.path)


class GoogleSheetsSheet:
    """Live Google Sheet via a service account (Sheets API v4).

    Setup: share the sheet with the service-account email, set
    GOOGLE_SPREADSHEET_ID + GOOGLE_CREDENTIALS_FILE, `pip install
    google-api-python-client google-auth`. Reads A4:T (header row 4 like
    the test workbook); write-back uses values.batchUpdate on Q-T.
    """

    SCOPES = ["https://www.googleapis.com/auth/spreadsheets"]

    def __init__(self, spreadsheet_id: str, credentials_file: str,
                 sheet_range: str = "Passport Skill Records!A4:T"):
        if not spreadsheet_id or not credentials_file:
            raise RuntimeError("SHEET_SOURCE=google requires GOOGLE_SPREADSHEET_ID and GOOGLE_CREDENTIALS_FILE")
        self.spreadsheet_id, self.sheet_range = spreadsheet_id, sheet_range
        self.credentials_file = credentials_file
        self._service = None

    def _api(self):
        if self._service is None:
            try:
                from google.oauth2.service_account import Credentials
                from googleapiclient.discovery import build
            except ImportError as e:
                raise RuntimeError("SHEET_SOURCE=google requires `pip install google-api-python-client google-auth`") from e
            creds = Credentials.from_service_account_file(self.credentials_file, scopes=self.SCOPES)
            self._service = build("sheets", "v4", credentials=creds)
        return self._service.spreadsheets()

    def fetch(self) -> list[dict]:
        resp = self._api().values().get(spreadsheetId=self.spreadsheet_id,
                                        range=self.sheet_range).execute()
        values = resp.get("values", [])
        if not values:
            return []
        header_row = 4  # A4:T range starts at the header
        rows = []
        for offset, values_row in enumerate(values[1:], start=1):
            padded = list(values_row) + [""] * (len(COLUMNS) - len(values_row))
            r = {k: v for k, v in zip(COLUMNS, padded)}
            r["_sheetRow"] = header_row + offset
            rows.append(r)
        return _clean(rows)

    def write_back(self, updates: list[dict]) -> None:
        data = []
        sheet_name = self.sheet_range.split("!")[0]
        for u in updates:
            row = u.get("_sheetRow")
            if not row:
                continue
            data.append({"range": f"{sheet_name}!Q{row}:T{row}",
                         "values": [[u["syncStatus"] or "", u["syncedVersion"] or "",
                                     u["syncedAt"] or "", u["syncNotes"] or ""]]})
        if data:
            self._api().values().batchUpdate(
                spreadsheetId=self.spreadsheet_id,
                body={"valueInputOption": "RAW", "data": data}).execute()


def make_sheet_source(settings) -> Optional[SheetSource]:
    if settings.sheet_source == "google":
        return GoogleSheetsSheet(settings.google_spreadsheet_id,
                                 settings.google_credentials_file, settings.google_sheet_range)
    if settings.sheet_source == "xlsx":
        return XlsxSheet(settings.xlsx_path, settings.sheet_writeback)
    return LocalJsonSheet(settings.seed_file, settings.sheet_writeback)
